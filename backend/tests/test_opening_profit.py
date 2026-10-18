"""Public onboarding and profit-carry contracts; synthetic SQLite only."""
from copy import deepcopy
from io import BytesIO
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Account, PortfolioOpening, PortfolioEvent, Deal
from app.services.portfolio import service as s
from app.services.imports import preview_import, confirm_import
from app.main import _openpyxl_template, _aggregate_positions

ASSET = dict(code="CN.SYNTH", name="Synthetic", market="CN", currency="CNY", asset_type="stock", unit="share")
T = "2026-01-01T12:00:00Z"


@pytest.fixture
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine, autoflush=False) as db:
        db.add_all([Account(account_id="a"), Account(account_id="b")]); db.commit()
        yield db
    engine.dispose()


def opening(db, **changes):
    data = dict(effective_at=T, cash={"CNY": "20000"}, historical_profit_cny="15000",
                positions=[dict(ASSET, quantity="100", average_cost="900", price="1000")])
    data.update(changes)
    s.save_opening(db, "a", data)
    return data


def event(db, kind, day, **data):
    return s.save_event(db, dict(account_id="a", kind=kind, effective_at=f"2026-01-{day:02}T12:00:00Z", data=data))


def test_carry_is_not_cash_cost_or_closed_trade_and_survives_full_exit(db):
    data = opening(db)
    c = s.overview(db)["current"]
    assert D(c["total_assets"]) == 120000 and D(c["cash_value"]) == 20000
    assert D(c["cumulative_profit"]) == 25000 and D(c["profit"]) == 0
    assert not s.closed_positions(db)["items"]
    event(db, "deposit", 2, amount="1000")
    event(db, "valuation", 3, prices=[dict(ASSET, price="1010")])
    event(db, "dividend", 4, **ASSET, amount="500")
    s.save_trade(db, dict(ASSET, account_id="a", effective_at="2026-01-05T12:00:00Z", quantity="100", price="1010", side="SELL"))
    c = s.overview(db)["current"]
    assert D(c["profit"]) == 1500 and D(c["cumulative_profit"]) == 26500
    assert D(s.closed_positions(db)["items"][0]["realized"]) == 11500
    event(db, "withdrawal", 6, amount=c["cash_value"])
    assert D(s.overview(db)["current"]["cumulative_profit"]) == 26500
    event(db, "deposit", 7, amount="1000")
    assert D(s.overview(db)["current"]["cumulative_profit"]) == 26500
    data["historical_profit_cny"] = "-3000"
    before = deepcopy(s.overview(db)["current"])
    s.save_opening(db, "a", data, correction=True)
    after = s.overview(db)["current"]
    assert D(after["cumulative_profit"]) == 8500
    assert after["cash"] == before["cash"] and after["positions"] == before["positions"]


def test_missing_is_not_zero_and_scope_transfer_baseline_are_independent(db):
    opening(db)
    s.save_opening(db, "b", dict(effective_at=T, positions=[], cash={"CNY":"0"}))
    assert s.overview(db)["current"]["cumulative_profit"] is None
    assert s.overview(db)["current"]["known_cumulative_profit"] == "25000"
    row = s.overview(db, "b")["openings"][0]
    row["historical_profit_cny"] = "0"
    s.save_opening(db, "b", row, correction=True)
    event(db, "transfer", 2, amount="1000", to_account_id="b", fees="2")
    assert D(s.overview(db)["current"]["cumulative_profit"]) == 24998
    event(db, "valuation", 3, performance_baseline="all", prices=[dict(ASSET, price="1010")])
    current = s.overview(db)["current"]
    assert D(current["profit"]) == 0 and D(current["cumulative_profit"]) == 25998


def test_unknown_cost_does_not_block_period_profit_or_become_known(db):
    opening(db, positions=[dict(ASSET, quantity="100", price="1000")])
    current = s.overview(db)["current"]
    assert current["cumulative_profit"] is None and current["profit"] == "0"
    assert current["account_profits"][0]["opening_holding_profit"] is None


def workbook(*, cost="900", carry="15000", cash="20000"):
    wb = Workbook(); wb.remove(wb.active)
    for name, headers, rows in [
        ("账户起点", ["统计起始时间*", "历史结转盈亏（CNY）", "收益说明"], [[T,carry,"user confirmed"]]),
        ("资金余额", ["币种*", "现金余额*"], [["CNY",cash]]),
        ("当前持仓", ["标的代码*", "市场*", "币种*", "数量*", "摊薄成本价", "价格"], [["CN.SYNTH","CN","CNY","100",cost,"1000"]]),
    ]:
        ws=wb.create_sheet(name); ws.append(headers)
        for row in rows: ws.append(row)
    out=BytesIO(); wb.save(out); return out.getvalue()


def test_excel_preview_atomic_confirmation_and_duplicate_correction_guard(db):
    preview = preview_import("excel", "opening.xlsx", workbook(), "a")
    assert preview["can_confirm"] and not db.get(PortfolioOpening,"a")
    checked = s.preview_opening(db,"a",preview["opening_payload"])
    assert D(checked["current"]["cumulative_profit"]) == 25000
    assert not db.get(PortfolioOpening,"a")
    assert confirm_import(db,preview)["inserted_count"] == 1
    assert confirm_import(db,preview)["status"] == "已导入"
    assert not list(db.scalars(select(Deal))) and not list(db.scalars(select(PortfolioEvent)))
    data=s.overview(db,"a")["openings"][0]; data["historical_profit_cny"]="16000"
    s.save_opening(db,"a",data,correction=True)
    with pytest.raises(ValueError,match="已有记账起点"):
        confirm_import(db,preview)


def test_excel_negative_cost_blank_carry_and_missing_cash(db):
    preview=preview_import("excel","x.xlsx",workbook(cost="-1",carry=None),"a")
    assert preview["can_confirm"]
    confirm_import(db,preview)
    assert s.overview(db)["current"]["cumulative_profit"] is None
    assert s.overview(db)["current"]["positions"][0]["average_cost"] == "-1"
    bad=preview_import("excel","x.xlsx",workbook(cash=None),"b")
    assert not bad["can_confirm"]
    with pytest.raises(ValueError): confirm_import(db,bad)
    assert not db.get(PortfolioOpening,"b")


def test_opening_preview_failure_rolls_back_and_keeps_estimate_metadata(db):
    data=opening(db)
    event(db,"withdrawal",2,amount="20000")
    data["cash"]={"CNY":"0"}
    with pytest.raises(ValueError): s.preview_opening(db,"a",data,correction=True)
    assert db.get(PortfolioOpening,"a").data["cash"] == {"CNY":"20000"}


def test_known_diluted_cost_is_available_in_legacy_detail(db):
    opening(db)
    merged=_aggregate_positions(s.projection_positions(db),120000,"CNY",{"CNY":1})[0]
    assert merged["cost_method"] == "diluted" and merged["average_cost"] == 900


def test_subscription_fee_changes_profit_once_and_cancel_preserves_loss(db):
    opening(db, positions=[], cash={"CNY":"1000"}, historical_profit_cny="0")
    pending=event(db,"subscription",2,amount="200",fees="2")
    c=s.overview(db)["current"]
    assert c["cash_value"] == "798" and c["cumulative_profit"] == "-2"
    event(db,"cancel_subscription",3,pending_id=f"event:{pending}")
    assert s.overview(db)["current"]["cash_value"] == "998"
    assert s.overview(db)["current"]["cumulative_profit"] == "-2"


def test_settlement_amount_survives_note_correction(db):
    opening(db,positions=[],cash={"CNY":"1000"})
    payload=dict(ASSET,account_id="a",effective_at=T,side="BUY",quantity="10",price="10",gross_amount="99.99",fees="1")
    rid=s.save_trade(db,payload)
    before=s.overview(db)["current"]
    s.save_trade(db,{**payload,"note":"corrected note"},rid)
    after=s.overview(db)["current"]
    assert after["cash_value"] == before["cash_value"] == "899.01"
    assert after["positions"][0]["cost"] == "100.99"


def test_exported_opening_template_is_readable_and_rejects_unfilled_input():
    # Exercise the actual download producer rather than maintaining a second template.
    from app.services.portfolio.opening_import import template_sheets
    content=_openpyxl_template(template_sheets()).getvalue()
    wb=load_workbook(BytesIO(content),read_only=True)
    assert wb.sheetnames == ["账户起点","资金余额","当前持仓"]
    wb.close()
    preview=preview_import("excel","template.xlsx",content,"a")
    assert not preview["can_confirm"] and preview["errors"]


def test_legacy_excel_reconciliation_accepts_confirmed_negative_diluted_cost(db, monkeypatch):
    import app.services.sync as sync
    for name in ("_create_profile","_create_review","_mark_freshness_states","refresh_trade_reviews"):
        monkeypatch.setattr(sync,name,lambda *a:None)
    wb=Workbook(); wb.active.title="持仓快照"
    wb.active.append(["标的代码*","数量*","当前价格*","摊薄成本","快照时间"])
    wb.active.append(["CN.SYNTH",10,10,-2,T])
    out=BytesIO();wb.save(out)
    preview=preview_import("excel","legacy.xlsx",out.getvalue(),"a")
    confirm_import(db,{**preview,"confirm_diluted_cost":True})
    from app.models import PositionSnapshot
    row=db.scalar(select(PositionSnapshot))
    assert row.cost_method == "diluted" and row.average_cost == -2


def test_backfill_checks_cny_basis_and_retains_original_carry_checkpoint(db):
    data=dict(effective_at="2026-01-05T12:00:00Z",cash={"CNY":"20000"},historical_profit_cny="15000",
              positions=[dict(ASSET,quantity="100",average_cost="900",price="1000")])
    s.save_opening(db,"a",data)
    payload=dict(opening=dict(effective_at=T,cash={"CNY":"110000"},historical_profit_cny="15000"),
        records=[dict(kind="trade",effective_at="2026-01-02T12:00:00Z",data=dict(ASSET,quantity="100",price="900",side="BUY"))])
    preview=s.rebuild_history(db,"a",payload,preview=True)
    assert not preview["differences"] and D(preview["current"]["cumulative_profit"]) == 25000
    s.rebuild_history(db,"a",payload)
    assert db.get(PortfolioOpening,"a").data["history_checkpoints"][0]["historical_profit_cny"] == "15000"
