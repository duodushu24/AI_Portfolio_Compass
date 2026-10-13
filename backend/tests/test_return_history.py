"""Phase I contracts, using synthetic holdings and a separate SQLite database."""
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.main import app
from app.models import Account, Deal, PositionSnapshot, ReturnObservation
from app.services.holding_history import new_batch, positions_as_of
from app.services.return_history import refresh_observations
from app.services.sync import latest_positions

T = datetime(2026, 10, 1, 8)


@pytest.fixture
def records(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'returns.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, autoflush=False)
    def isolated_db():
        with sessions() as db: yield db
    app.dependency_overrides[get_db] = isolated_db
    import app.services.sync as sync
    # Import tests exercise persistence; exclude unrelated profile/AI/market work.
    for name in ("_create_profile", "_create_review", "_mark_freshness_states", "refresh_trade_reviews"):
        monkeypatch.setattr(sync, name, lambda *args: None)
    with sessions() as db:
        db.add_all([Account(account_id="synthetic-a", base_currency="CNY", enabled=True),
                    Account(account_id="synthetic-b", base_currency="CNY", enabled=True)])
        db.commit()
    client = TestClient(app)  # No startup, real runtime/database/knowledge untouched.
    yield client, sessions
    client.close()
    app.dependency_overrides.pop(get_db, None)
    engine.dispose()


def position(db, code="CN.TEST", at=T, account="synthetic-a", quantity=100, average=10,
             currency="CNY", cost_method="average", batch=None, **kw):
    row = PositionSnapshot(account_id=account, code=code, name="虚构测试标的", market=code.split(".")[0],
        quantity=quantity, average_cost=average, current_price=12, raw_currency=currency,
        normalized_currency=currency, raw_market_value=quantity * 12, normalized_market_value=quantity * 12,
        cost_method=cost_method, snapshot_time=at, sync_id="synthetic", batch_id=batch.id if batch else None, **kw)
    db.add(row); db.flush()
    return row


def observe(client, *, at=T + timedelta(days=1), currency="CNY", account="all", **updates):
    draft = client.get("/api/returns/draft", params={"account_id": account, "observed_at": at.isoformat() + "Z", "currency": currency}).json()
    inputs = [{"position_id": i["position_id"], "account_id": i["account_id"], "code": i["code"],
        "price": "12", "price_at": at.isoformat() + "Z", "source": "虚构测试", "confirm_average_cost": True, **updates} for i in draft["items"]]
    payload = {"observed_at": at.isoformat() + "Z", "account_ids": draft["account_ids"], "currency": currency, "inputs": inputs}
    response = client.post("/api/returns/observations", json=payload)
    assert response.status_code == 201, response.text
    return response.json()["item"], payload


def test_full_partial_empty_scope_and_legacy_as_of(records):
    _, sessions = records
    with sessions() as db:
        legacy = position(db, "CN.OLD")
        legacy.market = "SH"  # Legacy/Futu exchange aliases belong to the CN full scope.
        position(db, "HK.KEEP", currency="HKD")
        position(db, "CN.OLD", account="synthetic-b")
        partial = new_batch(db, "synthetic-a", T + timedelta(days=1))
        position(db, "CN.NEW", T + timedelta(days=1), batch=partial)
        assert len(positions_as_of(db)) == 4
        full = new_batch(db, "synthetic-a", T + timedelta(days=2), completeness="full", market="CN")
        position(db, "CN.NEW", T + timedelta(days=2), batch=full)
        db.commit()
        assert {(p.account_id, p.code) for p in latest_positions(db)} == {
            ("synthetic-a", "CN.NEW"), ("synthetic-a", "HK.KEEP"), ("synthetic-b", "CN.OLD")}
        assert len(positions_as_of(db, at=T + timedelta(days=1))) == 4
        new_batch(db, "synthetic-a", T + timedelta(days=3), completeness="full")
        db.commit()
        assert [p.account_id for p in latest_positions(db)] == ["synthetic-b"]
        assert len(positions_as_of(db, at=T)) == 3  # Legacy completeness remains unknown.


def test_same_day_batches_have_deterministic_order(records):
    _, sessions = records
    with sessions() as db:
        b = new_batch(db, "synthetic-a", T, completeness="full")
        position(db, batch=b)
        new_batch(db, "synthetic-a", T, completeness="full")
        db.commit()
        assert positions_as_of(db) == []


def test_weighted_cross_account_curve_and_frozen_price(records):
    client, sessions = records
    with sessions() as db:
        position(db)
        position(db, account="synthetic-b", quantity=100, average=20)
        db.commit()
    item, _ = observe(client)
    assert Decimal(item["return_ratio"]) == Decimal("-0.2")  # 2400/3000 - 1, not mean of 20% and -40%.
    assert item["asset_returns"]["CN.TEST"] == "-0.2"
    with sessions() as db:
        p = db.scalar(select(PositionSnapshot).where(PositionSnapshot.account_id == "synthetic-a"))
        p.current_price = 999
        refresh_observations(db); db.commit()
    series = client.get("/api/returns/series").json()
    assert series["items"][0]["items"][0]["price"] == "12"
    assert series["items"][0]["return_ratio"] == "-0.2"
    assert series["cumulative_available"] is False
    assert "cumulative" not in series["items"][0]


def test_missing_cost_price_and_empty_are_not_zero_returns(records):
    client, sessions = records
    with sessions() as db:
        position(db, cost_method="unknown"); db.commit()
    unknown, _ = observe(client, confirm_average_cost=False)
    assert unknown["return_ratio"] is None and unknown["covered_count"] == 0
    missing, _ = observe(client, price=None)
    assert missing["return_ratio"] is None
    loss, _ = observe(client, price="0")
    assert loss["return_ratio"] == "-1"  # Explicit zero is a real price, unlike missing.
    empty = client.post("/api/returns/accounts/synthetic-a/empty", json={"observed_at": (T + timedelta(days=2)).isoformat() + "Z", "confirmation": "EMPTY_SCOPE"})
    assert empty.status_code == 200
    item, _ = observe(client, at=T + timedelta(days=3))
    assert item["status"] == "empty" and item["return_ratio"] is None
    assert client.get("/api/returns/history").json()["items"]


def test_foreign_cost_basis_is_not_today_fx(records):
    client, sessions = records
    with sessions() as db:
        position(db, "US.TEST", currency="USD", quantity=10); db.commit()
    unknown, _ = observe(client, fx_rate="8")
    assert unknown["return_ratio"] is None
    item, _ = observe(client, fx_rate="8", cost_cny_amount="700")
    assert item["items"][0]["market_value"] == "960.0"
    assert Decimal(item["return_ratio"]) == (Decimal(960) - 700) / 700
    original, _ = observe(client, currency="USD")
    assert Decimal(original["return_ratio"]) == Decimal(".2")


def test_new_state_correction_and_void_do_not_resurrect(records):
    client, sessions = records
    with sessions() as db: position(db); db.commit()
    payload = {"code": "CN.TEST", "quantity": 50, "average_cost": 10, "current_price": 12,
               "currency": "CNY", "market": "CN", "cost_method": "average", "snapshot_time": "2026-10-01T18:00:00+08:00"}
    saved = client.post("/api/data/accounts/synthetic-a/positions", json=payload)
    assert saved.status_code == 200, saved.text
    with sessions() as db:
        assert len(db.scalars(select(PositionSnapshot)).all()) == 2
        assert latest_positions(db, "synthetic-a")[0].snapshot_time == datetime(2026, 10, 1, 10)
        assert positions_as_of(db, at=T)[0].quantity == 100
    assert client.post("/api/data/accounts/synthetic-a/positions", json=payload).status_code == 409
    corrected = client.post("/api/data/accounts/synthetic-a/positions", json={**payload, "save_mode": "correction",
        "original_code": "CN.TEST", "original_snapshot_time": payload["snapshot_time"], "quantity": 60})
    assert corrected.status_code == 200
    deleted = client.delete("/api/data/accounts/synthetic-a/positions", params={"code": "CN.TEST", "snapshot_time": payload["snapshot_time"]})
    assert deleted.status_code == 200
    with sessions() as db:
        assert latest_positions(db, "synthetic-a") == []
        assert positions_as_of(db, at=T)[0].quantity == 100
        assert len(db.scalars(select(PositionSnapshot)).all()) == 2
    uncertain, _ = observe(client)
    assert uncertain["status"] == "partial" and uncertain["return_ratio"] is None
    assert uncertain["items"][0]["quantity"] is None  # Voiding a fact is not confirmed clearance.


def test_backfill_recalculates_quantity_without_fabricating_trade(records):
    client, sessions = records
    with sessions() as db: position(db); db.commit()
    earlier, _ = observe(client, at=T + timedelta(hours=1))
    later, _ = observe(client, at=T + timedelta(days=5))
    saved = client.post("/api/data/accounts/synthetic-a/positions", json={"code": "CN.TEST", "quantity": 50,
        "average_cost": 8, "cost_method": "average", "current_price": 10, "currency": "CNY", "market": "CN",
        "snapshot_time": (T + timedelta(days=2)).isoformat() + "Z"})
    assert saved.status_code == 200
    data = client.get("/api/returns/series").json()
    assert data["items"][0]["id"] == earlier["id"] and data["items"][0]["items"][0]["quantity"] == "100.0"
    assert data["items"][1]["id"] == later["id"] and data["items"][1]["items"][0]["quantity"] == "50.0"
    assert Decimal(data["items"][1]["return_ratio"]) == Decimal(".5")
    assert data["items"][1]["holdings_changed"]
    assert data["events"] == []


def test_known_deals_filtered_by_scope_and_time(records):
    client, sessions = records
    with sessions() as db:
        position(db)
        position(db, at=T + timedelta(hours=2), quantity=50)
        db.add(Deal(account_id="synthetic-a", deal_id="synthetic-sale", code="CN.TEST", side="SELL", price=12,
            quantity=50, deal_time=T + timedelta(hours=1), raw_payload={"fee": 1}))
        db.commit()
    observe(client, at=T + timedelta(hours=3))
    data = client.get("/api/returns/series", params={"account_id": "synthetic-a"}).json()
    assert data["items"] == []  # Full-account observation is not mixed into a single account curve.
    assert data["other_scope_count"] == 1
    event = data["events"][0]
    assert event["before"]["quantity"] == 100 and event["after"]["quantity"] == 50
    assert event["fees"]["fee"] == 1
    assert "realized_pnl" not in event
    assert client.get("/api/returns/series", params={"start": "2026-10-02T00:00:00Z"}).json()["events"] == []


def test_observation_edit_stale_missing_and_no_future_price(records):
    client, sessions = records
    with sessions() as db: position(db); db.commit()
    item, payload = observe(client, price_at=T.isoformat() + "Z", price_kind="nav")
    assert item["price_warning"] and item["items"][0]["stale"]
    payload["inputs"][0]["price"] = "13"
    changed = client.put(f"/api/returns/observations/{item['id']}", json=payload)
    assert changed.status_code == 200 and Decimal(changed.json()["item"]["return_ratio"]) == Decimal(".3")
    payload["inputs"][0]["price_at"] = "2026-10-09T00:00:00Z"
    assert client.post("/api/returns/observations", json=payload).status_code == 400
    payload["inputs"][0]["price_at"] = None
    assert client.post("/api/returns/observations", json=payload).status_code == 400


def test_excel_scope_and_failure_do_not_clear_account(records):
    client, sessions = records
    with sessions() as db: position(db, "CN.OLD"); db.commit()
    def preview(code):
        wb = Workbook(); wb.active.title = "持仓快照"
        wb.active.append(["标的代码*", "数量*", "当前价格*", "平均成本"])
        wb.active.append([code, 20, 12, 10])
        buf = BytesIO(); wb.save(buf)
        response = client.post("/api/import/excel/preview?account_id=synthetic-a", files={"file": ("synthetic.xlsx", buf.getvalue())})
        assert response.status_code == 200
        return response.json()
    partial = preview("CN.NEW")
    assert client.post("/api/import/excel/confirm", json=partial).json()["status"] == "成功"
    with sessions() as db: assert len(latest_positions(db, "synthetic-a")) == 2
    full = {**preview("CN.FULL"), "holding_completeness": "full", "holding_observed_at": "2026-10-05T08:00:00Z", "confirm_average_cost": True}
    # Full recorded in the past leaves a later partial intact; actual times govern, not entry order.
    result = client.post("/api/import/excel/confirm", json=full)
    assert result.json()["status"] == "成功", result.text
    with sessions() as db:
        assert [p.code for p in positions_as_of(db, "synthetic-a", datetime(2026, 10, 5, 9))] == ["CN.FULL"]
    invalid = {**preview("CN.BAD"), "holding_completeness": "full", "holding_market": "HK", "holding_observed_at": "2026-10-06T08:00:00Z"}
    assert client.post("/api/import/excel/confirm", json=invalid).status_code == 400
    with sessions() as db: assert "CN.NEW" in {p.code for p in latest_positions(db, "synthetic-a")}


def test_corrected_cost_cannot_reuse_old_foreign_confirmation(records):
    client, sessions = records
    with sessions() as db:
        row = position(db, "US.TEST", currency="USD", cost_method="unknown")
        db.commit()
    observe(client, fx_rate="7", cost_cny_amount="7000")
    corrected = client.post("/api/data/accounts/synthetic-a/positions", json={"save_mode": "correction",
        "original_code": "US.TEST", "original_snapshot_time": T.isoformat() + "Z", "code": "US.TEST",
        "snapshot_time": T.isoformat() + "Z", "quantity": 200, "average_cost": 10, "current_price": 12,
        "currency": "USD", "market": "US", "cost_method": "average"})
    assert corrected.status_code == 200
    item = client.get("/api/returns/series").json()["items"][0]
    assert item["return_ratio"] is None
    assert "缺少已确认的折算成本" in item["items"][0]["problems"]


def test_existing_sqlite_migration_keeps_legacy_observation(tmp_path, monkeypatch):
    from sqlalchemy import text
    from app import database
    engine = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE position_snapshots (id integer primary key, code varchar)"))
        conn.execute(text("INSERT INTO position_snapshots VALUES (1, 'SYNTHETIC.LEGACY')"))
    monkeypatch.setattr(database, "engine", engine)
    monkeypatch.setattr(database.settings, "database_url", f"sqlite:///{tmp_path / 'legacy.db'}")
    database.init_db()
    with engine.connect() as conn:
        legacy = conn.execute(text("SELECT code, batch_id, cost_method, is_void FROM position_snapshots")).one()
        assert legacy == ("SYNTHETIC.LEGACY", None, "unknown", 0)
        assert conn.execute(text("SELECT count(*) FROM return_observations")).scalar() == 0
    engine.dispose()


def test_reconfirming_import_with_full_scope_is_not_skipped(records):
    client, sessions = records
    with sessions() as db: position(db, "CN.OMITTED"); db.commit()
    wb = Workbook(); wb.active.title = "持仓快照"
    wb.active.append(["标的代码*", "数量*", "当前价格*", "平均成本", "快照时间"])
    wb.active.append(["CN.INCLUDED", 20, 12, 10, "2026-10-02T08:00:00Z"])
    buf = BytesIO(); wb.save(buf)
    preview = client.post("/api/import/excel/preview?account_id=synthetic-a", files={"file": ("synthetic.xlsx", buf.getvalue())}).json()
    assert client.post("/api/import/excel/confirm", json=preview).json()["status"] == "成功"
    with sessions() as db: assert len(latest_positions(db, "synthetic-a")) == 2
    full = {**preview, "holding_completeness": "full", "holding_observed_at": "2026-10-02T08:00:00Z", "confirm_average_cost": True}
    response = client.post("/api/import/excel/confirm", json=full)
    assert response.json()["status"] == "成功", response.text
    assert client.post("/api/import/excel/confirm", json=full).json()["status"] == "已导入"
    with sessions() as db:
        assert [p.code for p in latest_positions(db, "synthetic-a")] == ["CN.INCLUDED"]


def test_failed_full_import_rolls_back_scope_and_rows(records):
    client, sessions = records
    with sessions() as db:
        position(db, "CN.OLD")
        position(db, "CN.CONFLICT", at=T + timedelta(days=1)); db.commit()
    wb = Workbook(); wb.active.title = "持仓快照"
    wb.active.append(["标的代码*", "数量*", "当前价格*", "平均成本"])
    wb.active.append(["CN.NEW", 20, 12, 10])
    wb.active.append(["CN.CONFLICT", 20, 12, 10])
    buf = BytesIO(); wb.save(buf)
    preview = client.post("/api/import/excel/preview?account_id=synthetic-a", files={"file": ("synthetic.xlsx", buf.getvalue())}).json()
    full = {**preview, "holding_completeness": "full", "holding_observed_at": "2026-10-02T08:00:00Z"}
    result = client.post("/api/import/excel/confirm", json=full)
    assert result.json()["status"] == "失败", result.text
    with sessions() as db:
        assert {p.code for p in latest_positions(db, "synthetic-a")} == {"CN.OLD", "CN.CONFLICT"}


def test_old_account_scope_remains_queryable_after_enabling_another_account(records):
    client, sessions = records
    with sessions() as db: position(db); db.commit()
    saved, _ = observe(client)
    with sessions() as db:
        db.add(Account(account_id="synthetic-c", enabled=True)); db.commit()
    assert client.get("/api/returns/series").json()["items"] == []
    old = client.get("/api/returns/series", params=[("account_ids", "synthetic-a"), ("account_ids", "synthetic-b")]).json()
    assert old["items"][0]["id"] == saved["id"]
    assert ["synthetic-a", "synthetic-b"] in old["available_scopes"]


def test_currency_correction_does_not_reinterpret_frozen_quote(records):
    client, sessions = records
    with sessions() as db:
        position(db, "US.TEST", currency="USD", cost_cny_amount=7000); db.commit()
    observe(client, fx_rate="7")
    response = client.post("/api/data/accounts/synthetic-a/positions", json={"save_mode": "correction",
        "code": "US.TEST", "original_code": "US.TEST", "original_snapshot_time": T.isoformat() + "Z",
        "snapshot_time": T.isoformat() + "Z", "quantity": 100, "average_cost": 10, "cost_method": "average",
        "current_price": 12, "currency": "CNY", "market": "US"})
    assert response.status_code == 200
    item = client.get("/api/returns/series").json()["items"][0]
    assert item["return_ratio"] is None and item["items"][0]["price"] is None
    assert "报价币种与持仓不一致，请重新核对观察" in item["items"][0]["problems"]
