"""Standard account-opening workbook. No broker-specific formats or live data."""
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO
import json

from app.models import PortfolioOpening
from . import service
from .performance import instant, iso, public


def template_sheets():
    return [
        ("账户起点", ["统计起始时间*", "历史结转盈亏（CNY）", "收益说明"], [None] * 3),
        ("资金余额", ["币种*", "现金余额*", "兑人民币汇率", "汇率实际时间", "汇率来源", "汇率为估算"], [None] * 6),
        ("当前持仓", ["标的代码*", "标的名称", "市场*", "币种*", "资产类型", "数量*", "摊薄成本价", "人民币剩余成本", "价格", "持仓市值", "价格时间"], [None] * 11),
    ]


def signature(row):
    data = {k: v for k, v in row.data.items() if k != "opening_import"}
    return sha256(json.dumps(public([row.effective_at, row.included_at, row.excluded_at, data]), sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _rows(sheet):
    values = iter(sheet.values)
    headers = [str(v or "").strip() for v in next(values, [])]
    return [(i, dict(zip(headers, row))) for i, row in enumerate(values, 2) if any(v not in (None, "") for v in row)]


def _time(value):
    if value in (None, ""):
        raise ValueError("请填写统计起始时间")
    if not isinstance(value, datetime):
        value = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone(timedelta(hours=8)))
    return iso(value)


def preview(content, account_id, digest):
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if "账户起点" not in workbook.sheetnames:
            return None
        errors, warnings, positions, cash, fx = [], [], [], {}, {}
        payload = {"positions": positions, "cash": cash, "fx": fx, "source": "标准账户起点 Excel"}
        try:
            if not account_id:
                raise ValueError("请先在页面选择账户")
            if "资金余额" not in workbook.sheetnames or "当前持仓" not in workbook.sheetnames:
                raise ValueError("起点模板必须包含账户起点、资金余额、当前持仓三个工作表")
            rows = _rows(workbook["账户起点"])
            if len(rows) != 1:
                raise ValueError("账户起点需填写一行，一份文件只对应一个账户")
            row = rows[0][1]
            payload.update(effective_at=_time(row.get("统计起始时间*")), historical_profit_cny=row.get("历史结转盈亏（CNY）"),
                           historical_profit_note=str(row.get("收益说明") or ""))
            if payload["historical_profit_cny"] not in (None, ""):
                payload["historical_profit_cny"] = str(service.number(payload["historical_profit_cny"]))
            for line, row in _rows(workbook["资金余额"]):
                curr = service.currency(row.get("币种*"))
                if curr in cash:
                    raise ValueError(f"资金余额第 {line} 行币种重复")
                if row.get("现金余额*") in (None, ""):
                    raise ValueError(f"资金余额第 {line} 行现金未填写；无现金请明确填 0")
                cash[curr] = str(service.number(row["现金余额*"], minimum=service.ZERO))
                if row.get("兑人民币汇率") not in (None, ""):
                    estimate = row.get("汇率为估算")
                    if estimate not in (None, "", True, False, "是", "否"):
                        raise ValueError("汇率为估算请填写是或否")
                    fx[curr] = {"rate": row["兑人民币汇率"], "at": _time(row["汇率实际时间"]) if row.get("汇率实际时间") else payload["effective_at"],
                                "source": str(row.get("汇率来源") or "用户确认"), "is_estimate": estimate in (True, "是")}
            if not cash:
                raise ValueError("请填写至少一种币种现金，无现金请明确填 0")
            for line, row in _rows(workbook["当前持仓"]):
                if not str(row.get("标的代码*") or "").strip():
                    raise ValueError(f"当前持仓第 {line} 行缺少标的代码")
                item = service.asset({"code": row.get("标的代码*"), "name": row.get("标的名称"), "market": row.get("市场*"),
                    "currency": row.get("币种*"), "asset_type": row.get("资产类型") or "stock"})
                item.update(quantity=str(service.number(row.get("数量*"), minimum=service.ZERO)),
                    average_cost=None if row.get("摊薄成本价") in (None, "") else str(service.number(row["摊薄成本价"])),
                    cost_cny_amount=None if row.get("人民币剩余成本") in (None, "") else str(service.number(row["人民币剩余成本"])))
                price, value = row.get("价格"), row.get("持仓市值")
                if price not in (None, "") and value not in (None, ""):
                    raise ValueError(f"当前持仓第 {line} 行价格与市值只能填写一种")
                if value not in (None, ""):
                    quantity = service.number(item["quantity"])
                    if not quantity:
                        raise ValueError(f"当前持仓第 {line} 行市值换算需要非零数量")
                    price = service.number(value, minimum=service.ZERO) / quantity
                item.update(price=None if price in (None, "") else str(service.number(price, minimum=service.ZERO)),
                            price_at=_time(row["价格时间"]) if row.get("价格时间") else payload["effective_at"])
                positions.append(item)
            service.quotes({"prices": [{**p, "at": p["price_at"]} for p in positions], "fx": fx}, instant(payload["effective_at"]))
        except (ValueError, KeyError, TypeError) as exc:
            errors.append(str(exc))
        if payload.get("historical_profit_cny") in (None, ""):
            warnings.append("未提供历史结转盈亏，将仅展示可计算的统计期盈亏")
        return {"import_kind": "opening", "source_name": "excel", "import_hash": digest, "account": {"account_id": account_id},
                "opening_payload": public(payload), "positions": positions, "position_count": len(positions), "deals": [], "deal_count": 0,
                "total_assets": 0, "market_value": 0, "cash": 0, "errors": errors, "warnings": warnings, "can_confirm": not errors}
    finally:
        workbook.close()


def confirm(db, preview):
    aid = str(preview.get("account", {}).get("account_id", ""))
    digest = str(preview.get("import_hash", ""))
    old = db.get(PortfolioOpening, aid)
    if old:
        marker = old.data.get("opening_import", {})
        if marker.get("digest") == digest and marker.get("signature") == signature(old):
            return {"status": "已导入", "inserted_count": 0, "updated_count": 0}
        raise ValueError("账户已有记账起点；请使用起点纠错，持仓文件仅用于核对，不能覆盖账目")
    payload = deepcopy(preview["opening_payload"])
    def write():
        service.save_opening(db, aid, payload, commit=False)
        db.flush()
        row = db.get(PortfolioOpening, aid)
        row.data = {**row.data, "opening_import": {"digest": digest, "signature": signature(row)}}
        return {"status": "已建立账户起点", "inserted_count": 1, "updated_count": 0}
    return service.transaction(db, write)
